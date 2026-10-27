# The files the import tutorial opens: a Word manuscript, a second Word file to add to it and a zipped
# Scrivener project, in Italian and in English. python3 make.py writes them next to this script.
import os, zipfile
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))

BOOK = {
    'it': {
        'title': 'Il faro', 'author': 'Luca Di Gialleonardo',
        'parts': [
            ('Parte prima: Il paese', [
                ('Capitolo 1: Arrivo', [
                    ['Al molo', 'Il traghetto attraccò alle sei. Giulia fu l’ultima a scendere.', 'Sul molo non c’era nessuno ad aspettarla, come aveva immaginato.'],
                    ['La pensione', 'La pensione stava in cima alla salita, come le avevano detto.', 'La padrona le diede la chiave senza chiederle niente.'],
                ]),
                ('Capitolo 2: Il guardiano', [
                    ['Il silenzio', 'Il guardiano del faro non parlava con nessuno da anni.', 'In paese dicevano che parlasse soltanto con il mare.'],
                ]),
            ]),
            ('Parte seconda: Il mare', [
                ('Capitolo 3: La tempesta', [
                    ['Il vento', 'Il vento cominciò a soffiare dopo cena.', 'Le barche sbattevano contro il molo come porte lasciate aperte.', 'Nessuno uscì di casa.'],
                    ['La lanterna', 'Giulia salì al faro con una lanterna.', 'La porta era aperta.', 'Dentro, la luce girava ancora.'],
                ]),
                ('Capitolo 4: Il mattino', [
                    ['Le barche', 'Al mattino il porto era pieno di barche rotte.', 'Il faro, per la prima volta, era spento.'],
                ]),
            ]),
        ],
        'more': ('Parte terza: Il ritorno', [
            ('Capitolo 5: La lettera', [['La busta', 'Una lettera aspettava Giulia alla pensione.', 'Non c’era il nome di chi l’aveva scritta.']]),
        ]),
        'more_name': 'Il faro, terza parte',
        'scriv': ('La casa del sale', [
            ('Capitolo 1', [('La salina', 'Marta torna alla salina del nonno.', ['Marta arrivò alla salina con il sole già basso.']),
                            ('Il registro', 'Un vecchio registro con un nome cancellato.', ['Nel cassetto c’era un registro rilegato in tela.'])]),
            ('Capitolo 2', [('Il nome', 'Marta cerca chi ha cancellato il nome.', ['In paese nessuno ricordava quel nome.'])]),
        ]),
    },
    'en': {
        'title': 'The Lighthouse', 'author': 'Luca Di Gialleonardo',
        'parts': [
            ('Part One: The Village', [
                ('Chapter 1: Arrival', [
                    ['The quay', 'The ferry docked at six. Julia was the last to step off.', 'Nobody was waiting for her on the quay, as she had expected.'],
                    ['The guesthouse', 'The guesthouse stood at the top of the hill, as they had told her.', 'The landlady gave her the key and asked her nothing.'],
                ]),
                ('Chapter 2: The Keeper', [
                    ['Silence', 'The lighthouse keeper had not spoken to anyone in years.', 'In the village they said he talked only to the sea.'],
                ]),
            ]),
            ('Part Two: The Sea', [
                ('Chapter 3: The Storm', [
                    ['The wind', 'The wind began to blow after dinner.', 'The boats banged against the quay like doors left open.', 'Nobody left the house.'],
                    ['The lantern', 'Julia climbed to the lighthouse with a lantern.', 'The door was open.', 'Inside, the light was still turning.'],
                ]),
                ('Chapter 4: Morning', [
                    ['The boats', 'By morning the harbour was full of broken boats.', 'The lighthouse, for the first time, was dark.'],
                ]),
            ]),
        ],
        'more': ('Part Three: The Return', [
            ('Chapter 5: The Letter', [['The envelope', 'A letter was waiting for Julia at the guesthouse.', 'It did not say who had written it.']]),
        ]),
        'more_name': 'The Lighthouse, part three',
        'scriv': ('The Salt House', [
            ('Chapter 1', [('The salt pans', 'Marta goes back to her grandfather’s salt pans.', ['Marta reached the salt pans with the sun already low.']),
                           ('The ledger', 'An old ledger with a name crossed out.', ['In the drawer there was a ledger bound in cloth.'])]),
            ('Chapter 2', [('The name', 'Marta looks for whoever crossed out the name.', ['Nobody in the village remembered that name.'])]),
        ]),
    },
}


def word(parts, path, title, author):
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.author = author
    for p, (part, chapters) in enumerate(parts):
        numbered = p == 0  # the first part numbers its scenes, the second leaves a blank line between them
        doc.add_heading(part, level=1)
        for chapter, scenes in chapters:
            doc.add_heading(chapter, level=2)
            for i, (name, *lines) in enumerate(scenes):
                if numbered:
                    # each scene opens with its number on a line: Dante takes those for chapters until told otherwise
                    doc.add_paragraph(str(i + 1))
                elif i:
                    doc.add_paragraph('')  # scenes apart by a blank line
                for line in lines:
                    doc.add_paragraph(line)
    doc.save(path)


def rtf(text):
    body = ''.join(ch if ord(ch) < 128 else f"\\u{ord(ch)}?" for ch in text)
    return '{\\rtf1\\ansi\\ansicpg1252{\\fonttbl\\f0\\fnil Palatino;}\\pard\\f0\\fs24 ' + body + '\\par}'


def scrivener(book, path):
    title, chapters = book
    items, files, n = [], {}, 0
    for chapter, scenes in chapters:
        children = []
        for name, synopsis, lines in scenes:
            n += 1
            uid = f'S{n}'
            children.append(f'<BinderItem UUID="{uid}" Type="Text"><Title>{name}</Title></BinderItem>')
            files[f'Files/Data/{uid}/content.rtf'] = rtf('\\par\n'.join(lines))
            files[f'Files/Data/{uid}/synopsis.txt'] = synopsis + '\n'
        n += 1
        items.append(f'<BinderItem UUID="C{n}" Type="Folder"><Title>{chapter}</Title><Children>{"".join(children)}</Children></BinderItem>')
    scrivx = f'''<?xml version="1.0" encoding="UTF-8"?>
<ScrivenerProject Identifier="T" Version="2.0"><Binder>
<BinderItem UUID="D" Type="DraftFolder"><Title>Draft</Title><Children>{"".join(items)}</Children></BinderItem>
</Binder><ProjectProperties><ProjectTitle>{title}</ProjectTitle><FullName>Luca Di Gialleonardo</FullName></ProjectProperties></ScrivenerProject>'''
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr(f'{title}.scriv/{title}.scrivx', scrivx)
        for name, text in files.items():
            z.writestr(f'{title}.scriv/{name}', text)


for lang, b in BOOK.items():
    out = os.path.join(HERE, lang)
    os.makedirs(out, exist_ok=True)
    word(b['parts'], os.path.join(out, f"{b['title']}.docx"), b['title'], b['author'])
    word([b['more']], os.path.join(out, f"{b['more_name']}.docx"), b['title'], b['author'])
    scrivener(b['scriv'], os.path.join(out, f"{b['scriv'][0]}.zip"))
    print(lang, sorted(os.listdir(out)))
